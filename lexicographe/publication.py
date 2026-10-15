"""Publication : dictionnaire imprimé (PDF, Word), livret de syllabaire, export pour Bloom,
dossier pour Dictionary App Builder (application Android).

Le contenu est préparé une seule fois (`contenu`), puis mis en page par chaque format.
Règles de mise en page :
* vedettes dans l'ordre alphabétique déclaré du projet, regroupées sous l'initiale
  (graphème initial selon l'alphabet : « ŋg » est une lettre) ;
* deux colonnes ; en haut de chaque page, la première et la dernière vedette de la page
  (mots-guides) ; numéros de page ;
* seuls les statuts choisis sont publiés (par défaut : validé et à vérifier, pas les brouillons).
"""

import html
import os
import shutil
from dataclasses import dataclass, field
from datetime import date

from . import APP_NAME, exports, lexique
from .projet import nom_fichier_sur


@dataclass
class Options:
    titre: str = ""
    sous_titre: str = ""
    auteurs: str = ""
    statuts: tuple = ("validé", "à vérifier")
    exemples: bool = True
    index_inverse: bool = True
    format_page: str = "A4"          # « A4 » ou « A5 »


@dataclass
class SensImprime:
    numero: int
    categorie: str
    glose: str
    definition: str
    exemples: list[tuple[str, str]] = field(default_factory=list)


@dataclass
class EntreeImprimee:
    vedette: str
    tete: str                         # pour les mots-guides (sans exposant)
    prononciation: str = ""
    type_morph: str = ""
    sens: list[SensImprime] = field(default_factory=list)
    variantes: list[str] = field(default_factory=list)
    composants: list[str] = field(default_factory=list)
    renvois: list[str] = field(default_factory=list)
    entree_id: int | None = None      # pour retrouver les médias (site web)
    renvois_vedettes: list[tuple[str, str]] = field(default_factory=list)   # (type, vedette visée)


@dataclass
class Contenu:
    titre: str
    sous_titre: str
    auteurs: str
    langue_vern: str
    langue_analyse: str
    lettres: list[tuple[str, list[EntreeImprimee]]]
    index: list[tuple[str, list[str]]]

    @property
    def nb_entrees(self):
        return sum(len(e) for _, e in self.lettres)


def contenu(projet, options: Options):
    conn = projet.conn
    abreviations = {cid: (abrev or nom) for cid, nom, abrev, _ in lexique.categories(conn)}
    entrees = [e for e in exports.entrees_completes(projet) if e.statut in options.statuts]
    par_id = {e.id: e for e in entrees}
    ordre = lexique.ordre_du_projet(conn)
    relations = {}
    for source, cible, typ in conn.execute("SELECT source_id, cible_id, type FROM relations"):
        relations.setdefault(source, []).append((cible, typ))

    lettres, courante = [], None
    for e in entrees:
        imp = EntreeImprimee(vedette=e.vedette, tete=e.tete, prononciation=e.prononciation, entree_id=e.id,
                             type_morph=lexique.TYPES_MORPH.get(e.type_morph, "") if e.type_morph
                             not in ("", "stem", "root") else "")
        for n, s in enumerate(e.sens, start=1):
            if not (s.glose or s.definition):
                continue
            imp.sens.append(SensImprime(n, abreviations.get(s.categorie_id, ""), s.glose, s.definition,
                                        [(x.vernaculaire, x.traduction) for x in s.exemples if x.vernaculaire]
                                        if options.exemples else []))
        for i, sens in enumerate(imp.sens, start=1):
            sens.numero = i
        imp.variantes = [f"{v.forme} ({v.type})" for v in e.variantes if v.type != "allomorphe"]
        imp.composants = [par_id[c].vedette for c in e.composants if c in par_id]
        imp.renvois_vedettes = [(typ, par_id[c].vedette) for c, typ in relations.get(e.id, []) if c in par_id
                                and typ != "variante de"]
        imp.renvois = [f"{typ} : {vedette}" for typ, vedette in imp.renvois_vedettes]
        initiale = ordre.initiale(e.tete).upper() if e.tete else "?"
        if courante is None or courante[0] != initiale:
            courante = (initiale, [])
            lettres.append(courante)
        courante[1].append(imp)

    index = []
    if options.index_inverse:
        for mot, renvois in lexique.index_inverse(conn):
            vedettes = [v for eid, v, _, _ in renvois if eid in par_id]
            if vedettes:
                index.append((mot, vedettes))
    vern = projet.reglage("langue_vern_nom") or projet.reglage("langue_vern_code")
    analyse = projet.reglage("langue_analyse_nom") or projet.reglage("langue_analyse_code")
    return Contenu(options.titre or f"Dictionnaire {vern} – {analyse}", options.sous_titre, options.auteurs,
                   vern, analyse, lettres, index)


# -- HTML d'une entrée (pour la mise en page Qt) ---------------------------------------------

def _e(texte):
    return html.escape(texte or "")


def _point(texte):
    """Point final, sauf si le texte se termine déjà par une ponctuation."""
    return "" if (texte or "").rstrip().endswith((".", "!", "?", "…", "»", ")")) else "."


def html_entree(e: EntreeImprimee):
    morceaux = [f"<b>{_e(e.vedette)}</b>"]
    if e.prononciation:
        morceaux.append(f" [{_e(e.prononciation)}]")
    if e.type_morph:
        morceaux.append(f" <i>{_e(e.type_morph)}</i>")
    if e.variantes:
        morceaux.append(f" ({_e('; '.join(e.variantes))})")
    plusieurs = len(e.sens) > 1
    for s in e.sens:
        morceaux.append(f" <b>{s.numero}.</b>" if plusieurs else "")
        if s.categorie:
            morceaux.append(f" <i>{_e(s.categorie)}</i>")
        morceaux.append(f" {_e(s.glose)}")
        if s.definition and s.definition != s.glose:
            morceaux.append(f" — {_e(s.definition)}")
        dernier = s.definition if (s.definition and s.definition != s.glose) else s.glose
        for vern, trad in s.exemples:
            morceaux.append(f" <i>{_e(vern)}</i>" + (f" « {_e(trad)} »" if trad else ""))
            dernier = "»" if trad else vern
        morceaux.append(_point(dernier))
    if e.composants:
        morceaux.append(f" <span style='color:#555'>Composé de : {_e(' + '.join(e.composants))}.</span>")
    if e.renvois:
        morceaux.append(f" <span style='color:#555'>{_e('. '.join(e.renvois))}.</span>")
    return f"<p style='margin:0'>{''.join(morceaux)}</p>"


# -- PDF (mise en page par Qt, deux colonnes, mots-guides) ---------------------------------------

class _MiseEnPage:
    """Place des blocs HTML en deux colonnes, page après page."""

    def __init__(self, chemin, format_page, police, titre_courant):
        from PySide6.QtCore import QMarginsF
        from PySide6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter
        self.writer = QPdfWriter(chemin)
        self.writer.setResolution(300)
        taille = QPageSize(QPageSize.A5 if format_page == "A5" else QPageSize.A4)
        self.writer.setPageLayout(QPageLayout(taille, QPageLayout.Portrait, QMarginsF(15, 15, 15, 15),
                                              QPageLayout.Millimeter))
        self.writer.setTitle(titre_courant)
        self.writer.setCreator(APP_NAME)
        self.peintre = QPainter(self.writer)
        self.police = police
        r = self.writer.pageLayout().paintRectPixels(self.writer.resolution())
        self.largeur, self.hauteur = r.width(), r.height()
        mm = self.writer.resolution() / 25.4
        self.haut_contenu = 9 * mm            # place des mots-guides
        self.bas_contenu = self.hauteur - 8 * mm
        self.goutiere = 6 * mm
        self.largeur_colonne = (self.largeur - self.goutiere) / 2
        self.page = 0
        self.premiere = self.derniere = None
        self.colonne, self.y = 0, self.haut_contenu
        self.y_colonnes = self.haut_contenu      # haut des colonnes sur la page (sous un titre de lettre)
        self.titre_courant = titre_courant

    def _document(self, html_bloc, largeur):
        from PySide6.QtGui import QTextDocument
        doc = QTextDocument()
        doc.documentLayout().setPaintDevice(self.writer)
        doc.setDefaultFont(self.police)
        doc.setDocumentMargin(0)
        doc.setHtml(html_bloc)
        doc.setTextWidth(largeur)
        return doc

    def _dessiner(self, doc, x, y):
        self.peintre.save()
        self.peintre.translate(x, y)
        doc.drawContents(self.peintre)
        self.peintre.restore()

    def _fin_de_page(self):
        """Mots-guides et numéro de la page qui se termine."""
        from PySide6.QtCore import QRectF, Qt
        self.page += 1
        self.peintre.save()
        police = self.peintre.font()
        police.setFamily(self.police.family())
        # taille d'en-tête fixe : avec la grande police d'un livret, le trait barrait le titre
        police.setPointSizeF(min(self.police.pointSizeF(), 9.5))
        self.peintre.setFont(police)
        haut = QRectF(0, 0, self.largeur, self.haut_contenu * 0.8)
        if self.premiere:
            police.setBold(True)
            self.peintre.setFont(police)
            self.peintre.drawText(haut, Qt.AlignLeft | Qt.AlignTop, self.premiere)
            self.peintre.drawText(haut, Qt.AlignRight | Qt.AlignTop, self.derniere)
            police.setBold(False)
            self.peintre.setFont(police)
        self.peintre.drawText(haut, Qt.AlignHCenter | Qt.AlignTop, self.titre_courant)
        self.peintre.drawLine(0, int(self.haut_contenu * 0.75), int(self.largeur), int(self.haut_contenu * 0.75))
        self.peintre.drawText(QRectF(0, self.bas_contenu, self.largeur, self.hauteur - self.bas_contenu),
                              Qt.AlignHCenter | Qt.AlignBottom, str(self.page))
        self.peintre.restore()
        self.premiere = self.derniere = None

    def nouvelle_page(self, fin=True):
        if fin:
            self._fin_de_page()
        self.writer.newPage()
        self.colonne, self.y = 0, self.haut_contenu
        self.y_colonnes = self.haut_contenu

    def pleine_largeur(self, html_bloc, espace_apres=0.0):
        """Bloc sur toute la largeur (titre de lettre) : commence sous le contenu des deux colonnes."""
        doc = self._document(html_bloc, self.largeur)
        h = doc.size().height()
        page_vide = self.colonne == 0 and self.y <= self.haut_contenu + 1
        # un titre de lettre ne reste pas seul en bas de page ; jamais de page blanche
        if not page_vide and (self.colonne == 1 or self.y + h > self.bas_contenu - 3 * h):
            self.nouvelle_page()
        self._dessiner(doc, 0, self.y)
        self.y += h + espace_apres
        self.y_colonnes = self.y

    def bloc(self, html_bloc, guide=None):
        doc = self._document(html_bloc, self.largeur_colonne)
        h = doc.size().height()
        if self.y + h > self.bas_contenu:
            if self.colonne == 0:
                self.colonne = 1
                self.y = self.y_colonnes
            else:
                self.nouvelle_page()
        x = self.colonne * (self.largeur_colonne + self.goutiere)
        self._dessiner(doc, x, self.y)
        self.y += h + self.police.pointSizeF() * self.writer.resolution() / 72 * 0.35
        if guide:
            self.premiere = self.premiere or guide
            self.derniere = guide

    def page_de_titre(self, html_bloc):
        doc = self._document(html_bloc, self.largeur)
        self._dessiner(doc, 0, self.hauteur * 0.3)
        self.writer.newPage()

    def terminer(self):
        self._fin_de_page()
        self.peintre.end()


def _police(projet, taille):
    from PySide6.QtGui import QFont
    nom = projet.reglage("police_vern") or "Charis SIL"
    police = QFont(nom)
    police.setPointSizeF(taille)
    return police


def _page_de_titre_html(c: Contenu):
    lignes = [f"<p align='center' style='font-size:26pt'><b>{_e(c.titre)}</b></p>"]
    if c.sous_titre:
        lignes.append(f"<p align='center' style='font-size:14pt'>{_e(c.sous_titre)}</p>")
    lignes.append(f"<p align='center' style='font-size:11pt'>{c.nb_entrees} entrées</p>")
    if c.auteurs:
        lignes.append(f"<p align='center' style='font-size:12pt'><br>{_e(c.auteurs)}</p>")
    lignes.append(f"<p align='center' style='font-size:10pt; color:#666'><br>{date.today().strftime('%Y')}"
                  f" — préparé avec {APP_NAME}</p>")
    return "".join(lignes)


def exporter_pdf(projet, chemin, options: Options):
    """Dictionnaire en PDF, deux colonnes. Retourne le nombre d'entrées."""
    c = contenu(projet, options)
    taille = 9.5 if options.format_page == "A4" else 8.5
    mp = _MiseEnPage(chemin, options.format_page, _police(projet, taille), c.titre)
    mp.page_de_titre(_page_de_titre_html(c))
    for lettre, entrees in c.lettres:
        mp.pleine_largeur(f"<p align='center' style='font-size:{taille * 2.2}pt'><b>{_e(lettre)}</b></p>",
                          espace_apres=mp.haut_contenu * 0.2)
        for e in entrees:
            mp.bloc(html_entree(e), guide=e.tete)
    if c.index:
        mp.nouvelle_page()
        mp.pleine_largeur(f"<p align='center' style='font-size:{taille * 1.8}pt'><b>Index "
                          f"{_e(c.langue_analyse)} – {_e(c.langue_vern)}</b></p>", espace_apres=mp.haut_contenu * 0.2)
        for mot, vedettes in c.index:
            mp.bloc(f"<p style='margin:0'><b>{_e(mot)}</b> : {_e(', '.join(vedettes))}</p>", guide=mot)
    mp.terminer()
    return c.nb_entrees


# -- Word (.docx) --------------------------------------------------------------------------------

def _champ(paragraphe, instruction):
    """Champ Word (STYLEREF, PAGE…) recalculé par Word à l'affichage."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    champ = OxmlElement("w:fldSimple")
    champ.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    texte = OxmlElement("w:t")
    texte.text = "…"
    run.append(texte)
    champ.append(run)
    paragraphe._p.append(champ)


def _deux_colonnes(section):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    sect = section._sectPr
    cols = sect.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "340")


def _paragraphe(doc, style):
    """Paragraphe avec style posé directement par son identifiant : python-docx, sinon, reparcourt
    tous les styles à chaque paragraphe (l'export de 3 000 entrées prenait plus de 20 s)."""
    p = doc.add_paragraph()
    p._p.style = style.style_id
    return p


def _run_style(paragraphe, texte, style):
    r = paragraphe.add_run(texte)
    r._r.style = style.style_id
    return r


def exporter_docx(projet, chemin, options: Options):
    """Dictionnaire en Word : deux colonnes, mots-guides calculés par Word (champs STYLEREF)."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    c = contenu(projet, options)
    doc = Document()
    police = projet.reglage("police_vern") or "Charis SIL"
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = police, Pt(9.5 if options.format_page == "A4" else 8.5)
    vedette = doc.styles.add_style("Vedette", WD_STYLE_TYPE.CHARACTER)
    vedette.font.bold = True
    entree_style = doc.styles.add_style("Entrée", WD_STYLE_TYPE.PARAGRAPH)
    entree_style.base_style = normal
    entree_style.paragraph_format.left_indent = Mm(4)
    entree_style.paragraph_format.first_line_indent = Mm(-4)
    entree_style.paragraph_format.space_after = Pt(2)
    lettre_style = doc.styles["Heading 1"]
    lettre_style.font.name, lettre_style.font.color.rgb = police, RGBColor(0x1F, 0x5B, 0x57)

    section = doc.sections[0]
    if options.format_page == "A5":
        section.page_width, section.page_height = Mm(148), Mm(210)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    for marge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, marge, Mm(15))

    # page de titre
    for texte, taille, gras in ((c.titre, 26, True), (c.sous_titre, 14, False), (f"{c.nb_entrees} entrées", 11, False),
                                (c.auteurs, 12, False), (f"{date.today().year} — préparé avec {APP_NAME}", 10, False)):
        if texte:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(texte)
            r.font.size, r.font.bold = Pt(taille), gras

    corps = doc.add_section(WD_SECTION.NEW_PAGE)
    _deux_colonnes(corps)
    corps.header.is_linked_to_previous = False
    entete = corps.header.paragraphs[0]
    _champ(entete, 'STYLEREF "Vedette"')
    entete.add_run("\t")
    entete.add_run(c.titre).italic = True
    entete.add_run("\t")
    _champ(entete, 'STYLEREF "Vedette" \\l')
    corps.footer.is_linked_to_previous = False
    pied = corps.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _champ(pied, "PAGE")

    for lettre, entrees in c.lettres:
        h = doc.add_paragraph(lettre, style="Heading 1")
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for e in entrees:
            p = _paragraphe(doc, entree_style)
            _run_style(p, e.vedette, vedette)
            if e.prononciation:
                p.add_run(f" [{e.prononciation}]")
            if e.type_morph:
                p.add_run(f" {e.type_morph}").italic = True
            if e.variantes:
                p.add_run(f" ({'; '.join(e.variantes)})")
            plusieurs = len(e.sens) > 1
            for s in e.sens:
                if plusieurs:
                    p.add_run(f" {s.numero}.").bold = True
                if s.categorie:
                    p.add_run(f" {s.categorie}").italic = True
                p.add_run(f" {s.glose}")
                if s.definition and s.definition != s.glose:
                    p.add_run(f" — {s.definition}")
                dernier = s.definition if (s.definition and s.definition != s.glose) else s.glose
                for vern, trad in s.exemples:
                    p.add_run(f" {vern}").italic = True
                    dernier = vern
                    if trad:
                        p.add_run(f" « {trad} »")
                        dernier = "»"
                p.add_run(_point(dernier))
            for texte in ([f"Composé de : {' + '.join(e.composants)}."] if e.composants else []) + \
                         ([". ".join(e.renvois) + "."] if e.renvois else []):
                r = p.add_run(" " + texte)
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if c.index:
        doc.add_section(WD_SECTION.NEW_PAGE)
        h = doc.add_paragraph(f"Index {c.langue_analyse} – {c.langue_vern}", style="Heading 1")
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for mot, vedettes in c.index:
            p = _paragraphe(doc, entree_style)
            _run_style(p, mot, vedette)
            p.add_run(" : " + ", ".join(vedettes))
    doc.save(chemin)
    return c.nb_entrees


# -- Syllabaire : livret PDF et export pour Bloom --------------------------------------------------

def exporter_livret_syllabaire(projet, lecons, chemin, mots_par_page=24, format_page="A4"):
    """Une leçon par page : la lettre en grand, les mots, les phrases."""
    from . import syllabaire as sy
    mp = _MiseEnPage(chemin, format_page, _police(projet, 16), f"Syllabaire — {projet.reglage('langue_vern_nom')}")
    titre = (f"<p align='center' style='font-size:30pt'><b>Syllabaire</b></p>"
             f"<p align='center' style='font-size:16pt'>{_e(projet.reglage('langue_vern_nom') or projet.nom)}</p>"
             f"<p align='center' style='font-size:12pt'>{len(lecons)} leçons</p>")
    mp.page_de_titre(titre)
    for i, l in enumerate(lecons):
        g = sy.affichage(l.grapheme)
        majuscule = g.upper() if g.upper() != g else ""
        mots = "  ".join(_e(m) for m, _ in l.mots[:mots_par_page])
        phrases = "".join(f"<p style='font-size:15pt'>{_e(t)}</p>" for _, t in l.phrases[:6])
        mp.pleine_largeur(
            f"<p align='center' style='font-size:14pt; color:#666'>Leçon {l.numero}</p>"
            f"<p align='center' style='font-size:72pt'><b>{_e(majuscule)} {_e(g)}</b></p>"
            f"<p align='center' style='font-size:22pt; line-height:150%'>{mots}</p>{phrases}")
        if i < len(lecons) - 1:
            mp.nouvelle_page()
    mp.terminer()
    return len(lecons)


def exporter_bloom(projet, lecons, dossier):
    """Fichiers pour l'outil « Decodable Reader » de Bloom : lettres de chaque étape (une étape
    par ligne, nouvelles lettres seulement) et liste de mots. Retourne les chemins créés."""
    from . import syllabaire as sy
    os.makedirs(dossier, exist_ok=True)
    etapes = os.path.join(dossier, "etapes-bloom.txt")
    with open(etapes, "w", encoding="utf-8") as f:
        f.write("\n".join(sy.affichage(l.grapheme) for l in lecons) + "\n")
    mots = os.path.join(dossier, "mots-bloom.txt")
    with open(mots, "w", encoding="utf-8") as f:
        f.write(" ".join(m for l in lecons for m, _ in l.mots) + "\n")
    lisez = os.path.join(dossier, "LISEZ-MOI-Bloom.txt")
    with open(lisez, "w", encoding="utf-8") as f:
        f.write("UTILISER CE SYLLABAIRE DANS BLOOM\n"
                "=================================\n\n"
                "1. Dans Bloom, ouvre une collection dans la langue, puis un livre « Decodable Reader ».\n"
                "2. Outil « Decodable Reader » > « Set up Stages » (configurer les étapes).\n"
                "3. Onglet « Letters » : ajoute une étape par ligne du fichier etapes-bloom.txt,\n"
                "   en donnant à chaque étape la lettre de la ligne correspondante.\n"
                "4. Onglet « Words » : colle le contenu de mots-bloom.txt (ou choisis ce fichier).\n"
                "Bloom indiquera alors, page après page, les mots lisibles à chaque étape.\n")
    return [etapes, mots, lisez]


# -- Dictionary App Builder (application Android) -------------------------------------------------

def preparer_dab(projet, dossier):
    """Dossier prêt à importer dans Dictionary App Builder : fichier LIFT, sons (audio/) et images
    (pictures/) référencés depuis le LIFT, et mode d'emploi. Retourne le chemin du LIFT."""
    nom = nom_fichier_sur(projet.nom)
    racine = os.path.join(dossier, f"{nom}-DictionaryAppBuilder")
    os.makedirs(os.path.join(racine, "audio"), exist_ok=True)
    os.makedirs(os.path.join(racine, "pictures"), exist_ok=True)
    medias = {}
    for eid, sens_id, role, chemin in projet.conn.execute(
            """SELECT l.entree_id, l.sens_id, l.role, m.chemin FROM medias_liens l
               JOIN medias m ON m.id = l.media_id"""):
        if not os.path.isfile(chemin):
            continue
        sous = "audio" if role == "prononciation" else "pictures"
        cible = os.path.join(racine, sous, os.path.basename(chemin))
        if not os.path.exists(cible):
            shutil.copy2(chemin, cible)
        medias.setdefault(eid, []).append((role, sens_id, os.path.basename(chemin)))
    arbre, n = exports.construire_lift(projet, medias=medias)
    chemin_lift = os.path.join(racine, f"{nom}.lift")
    arbre.write(chemin_lift, encoding="utf-8", xml_declaration=True)
    with open(os.path.join(racine, "LISEZ-MOI.txt"), "w", encoding="utf-8") as f:
        f.write(f"APPLICATION ANDROID DU DICTIONNAIRE — {projet.nom}\n"
                "=========================================\n\n"
                "Ce dossier contient le dictionnaire au format LIFT, avec ses sons (audio/) et ses images\n"
                "(pictures/). Pour en faire une application Android :\n\n"
                "1. Installe Dictionary App Builder (gratuit, SIL) : https://software.sil.org/dictionaryappbuilder/\n"
                "2. Lance-le : « Nouveau projet » (New App), puis choisis le fichier "
                f"{nom}.lift de ce dossier.\n"
                "3. Les sons et les images sont trouvés automatiquement dans audio/ et pictures/.\n"
                "4. Règle le nom de l'application, les couleurs et l'icône, puis « Construire l'APK » (Build).\n\n"
                f"{n} entrées. Préparé avec {APP_NAME}.\n")
    return chemin_lift
