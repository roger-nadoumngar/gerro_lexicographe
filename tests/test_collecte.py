import os

import pytest

from lexicographe import collecte, domaines, fusion, lexique
from lexicographe.lexique import Entree, Exemple, Sens
from lexicographe.projet import Projet

DOMS = [
    {"guid": "d-corps", "code": "2.1", "parent": None, "noms": {"fr": "Corps"}, "descriptions": {"fr": "Le corps."},
     "questions": [{"question": {"fr": "(1) Les parties du corps ?"}, "mots": {"fr": "tête, main"}, "phrases": {}},
                   {"question": {"fr": "(2) Les os ?"}, "mots": {"fr": "os"}, "phrases": {}}]},
    {"guid": "d-ciel", "code": "1.1", "parent": None, "noms": {"fr": "Ciel"}, "descriptions": {}, "questions": []},
]


@pytest.fixture
def avec_domaines(projet):
    domaines.importer(projet.conn, DOMS)
    return projet


def test_questions_dans_l_ordre(avec_domaines):
    qs = collecte.questions(avec_domaines.conn)
    assert [(q.code, q.numero, q.total) for q in qs] == [("1.1", 0, 1), ("2.1", 0, 2), ("2.1", 1, 2)]
    assert qs[0].texte == "Quels mots se rapportent à « Ciel » ?", "question générique pour un domaine sans question"
    assert qs[1].mots_exemples == "tête, main"


def test_ajouter_mot_sans_doublon(avec_domaines):
    conn = avec_domaines.conn
    eid, cree = collecte.ajouter_mot(conn, "d-corps", "dɔ", "tête")
    assert cree
    e = lexique.lire_entree(conn, eid)
    assert (e.statut, e.sens[0].glose, e.sens[0].domaines) == ("brouillon", "tête", ["d-corps"])
    eid2, cree = collecte.ajouter_mot(conn, "d-ciel", "Dɔ")
    assert eid2 == eid and not cree, "même forme : le domaine est ajouté, pas de doublon"
    assert set(lexique.lire_entree(conn, eid).sens[0].domaines) == {"d-corps", "d-ciel"}
    assert collecte.mots_du_domaine(conn, "d-corps") == [(eid, "dɔ", "tête")]
    with pytest.raises(ValueError):
        collecte.ajouter_mot(conn, "d-corps", "  ")


def test_taches(avec_domaines, tmp_path):
    conn = avec_domaines.conn
    a = lexique.enregistrer_entree(conn, Entree(forme="ala", sens=[Sens(glose="")]))
    b = lexique.enregistrer_entree(conn, Entree(forme="ba", sens=[Sens(glose="père")]))
    c = lexique.enregistrer_entree(conn, Entree(forme="ka", sens=[Sens(glose="x", domaines=["d-corps"],
                                                                       exemples=[Exemple("ka ɓa", "x")])]))
    assert collecte.compter(conn) == {"glose": 1, "exemple": 1, "domaine": 2, "audio": 3}
    assert collecte.a_faire(conn, "glose") == [a]
    assert collecte.completer(conn, a, "glose", "dieu")
    assert lexique.lire_entree(conn, a).sens[0].glose == "dieu"
    assert not collecte.completer(conn, b, "glose", "autre"), "rien à compléter : rien n'est modifié"
    assert collecte.completer(conn, b, "exemple", ("ba ɓa", "le père vient"))
    assert lexique.lire_entree(conn, b).sens[0].exemples[0].traduction == "le père vient"
    assert collecte.completer(conn, a, "domaine", "d-ciel")
    lexique.lier_media(conn, c, str(tmp_path / "ka.wav"), "prononciation")
    assert collecte.compter(conn) == {"glose": 0, "exemple": 1, "domaine": 1, "audio": 2}


def test_position_de_reprise(projet):
    collecte.enregistrer_position(projet, "d-corps", 1)
    assert collecte.position_enregistree(projet) == {"guid": "d-corps", "numero": 1}


# -- fusion entre postes -------------------------------------------------------------------------

def _copie_poste(projet, tmp_path, nom):
    chemin = str(tmp_path / f"{nom}.lexique")
    projet.sauvegarder_vers(chemin)
    return Projet(chemin)


def test_fusion(projet, tmp_path):
    conn = projet.conn
    nom = lexique.categorie_par_nom(conn, "nom")
    commun = lexique.enregistrer_entree(conn, Entree(forme="ala", sens=[Sens(glose="dieu")]))
    garde = lexique.enregistrer_entree(conn, Entree(forme="ba", sens=[Sens(glose="père")]))
    poste = _copie_poste(projet, tmp_path, "poste2")
    # sur l'autre poste : une entrée modifiée, une nouvelle (avec catégorie, audio, composant), une phrase
    e = lexique.lire_entree(poste.conn, commun)
    e.sens[0].glose = "Dieu (créateur)"
    lexique.enregistrer_entree(poste.conn, e)
    poste.conn.execute("UPDATE entrees SET date_modification='2099-01-01T00:00:00' WHERE id=?", (commun,))
    poste.conn.execute("UPDATE entrees SET date_modification='2000-01-01T00:00:00' WHERE id=?", (garde,))
    poste.conn.commit()
    audio = tmp_path / "ko.wav"
    audio.write_bytes(b"RIFF")
    nouveau = lexique.enregistrer_entree(poste.conn, Entree(
        forme="kó-ala", composants=[commun], sens=[Sens(glose="arbre de dieu",
                                                         categorie_id=lexique.categorie_par_nom(poste.conn, "nom"))]))
    lexique.lier_media(poste.conn, nouveau, str(audio), "prononciation", "locuteur 2")
    with poste.conn:
        poste.conn.execute("INSERT INTO phrases (ref, vernaculaire, analyse, source) VALUES ('T009 1', 'kó', 'arbre', 'p2')")
    poste.fermer()
    avant = os.path.getmtime(str(tmp_path / "poste2.lexique"))

    r = fusion.fusionner(projet, str(tmp_path / "poste2.lexique"))
    assert (r.ajoutees, r.mises_a_jour, r.gardees, r.identiques, r.phrases, r.medias) == (1, 1, 1, 0, 1, 1)
    assert lexique.lire_entree(conn, commun).sens[0].glose == "Dieu (créateur)", "la version la plus récente gagne"
    assert lexique.lire_entree(conn, garde).sens[0].glose == "père"
    (eid,) = lexique.chercher_par_forme(conn, "kó-ala")
    ajoutee = lexique.lire_entree(conn, eid)
    assert ajoutee.composants == [commun] and ajoutee.sens[0].categorie_id == nom
    assert lexique.medias_de(conn, eid)[0][6] == "locuteur 2"
    assert os.path.getmtime(str(tmp_path / "poste2.lexique")) == avant, "l'autre fichier n'est pas modifié"

    r2 = fusion.fusionner(projet, str(tmp_path / "poste2.lexique"))
    assert (r2.ajoutees, r2.mises_a_jour, r2.phrases) == (0, 0, 0), "refusionner ne refait rien"


def test_fusion_refuse_le_projet_lui_meme(projet):
    with pytest.raises(ValueError):
        fusion.fusionner(projet, projet.chemin)


# -- collecte à partir d'une liste de mots ------------------------------------------------------

def _liste_word(chemin):
    import docx
    d = docx.Document()
    d.add_paragraph("ID\tFrançais\tNgambay")
    d.add_paragraph("1\teau\t")
    d.add_paragraph("2\tfeu\t")
    d.add_paragraph("1. Mammifères (2)")
    t = d.add_table(rows=3, cols=3)
    for i, (n, mot) in enumerate([("N°", "Nom commun"), ("1", "Lion"), ("2", "Eau")]):
        t.cell(i, 0).text, t.cell(i, 1).text = n, mot
    d.save(chemin)


def test_lire_une_liste(tmp_path):
    chemin = str(tmp_path / "liste.docx")
    _liste_word(chemin)
    assert [(e.mot, e.groupe) for e in collecte.lire_liste(chemin)] == [
        ("eau", ""), ("feu", ""), ("Lion", "Mammifères")], "en-têtes ignorés, doublon « Eau » retiré"
    texte = tmp_path / "liste.txt"
    texte.write_text("Français;Anglais\n1;main;hand\npied;foot\n", encoding="utf-8")
    assert [e.mot for e in collecte.lire_liste(str(texte))] == ["main", "pied"]


def test_segments_de_glose():
    assert collecte.segments_glose("bovin; boeuf, vache") == {"bovin", "boeuf", "vache"}
    assert collecte.segments_glose("piquer (scorpion, fourmi, etc.)") == {"piquer"}


def test_repondre_a_une_liste(projet):
    conn = projet.conn
    feu = lexique.enregistrer_entree(conn, Entree(forme="pə̀r", sens=[Sens(glose="feu")]))
    elements = [collecte.ElementListe("eau"), collecte.ElementListe("feu"), collecte.ElementListe("lion", "Mammifères")]
    collecte.enregistrer_liste(projet, "Swadesh", elements)
    assert collecte.listes(projet)["Swadesh"][2].groupe == "Mammifères"
    assert collecte.avancement(conn, elements) == 1
    assert collecte.trouves(conn, "Feu") == [(feu, "pə̀r")]
    eid, cree = collecte.repondre_liste(conn, elements[0], "màn̄", "Swadesh")
    assert cree and "Swadesh" in lexique.lire_entree(conn, eid).notes
    assert collecte.repondre_liste(conn, elements[0], "Màn̄", "Swadesh") == (eid, False), "pas de doublon"
    assert [s.glose for s in lexique.lire_entree(conn, eid).sens] == ["eau"]
    collecte.repondre_liste(conn, elements[2], "pə̀r", "Swadesh")
    assert [s.glose for s in lexique.lire_entree(conn, feu).sens] == ["feu", "lion"], "nouveau sens, même entrée"
    assert collecte.avancement(conn, elements) == 3
    collecte.enregistrer_position_liste(projet, "Swadesh", 2)
    assert collecte.position_liste(projet, "Swadesh") == 2
    collecte.supprimer_liste(projet, "Swadesh")
    assert collecte.listes(projet) == {}
