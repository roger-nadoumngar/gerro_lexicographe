import pytest

from lexicographe import interlineaire as il
from lexicographe import lexique, textes, usfm
from lexicographe.lexique import Entree, Sens, Variante
from lexicographe.morphologie import Analyseur


def _e(conn, forme, glose, type_morph="", variantes=()):
    return lexique.enregistrer_entree(conn, Entree(forme=forme, type_morph=type_morph, sens=[Sens(glose=glose)],
                                                   variantes=list(variantes)))


@pytest.fixture
def lexique_morpho(projet):
    conn = projet.conn
    ids = {
        "kó": _e(conn, "kó", "arbre", "stem"),
        "ɓe": _e(conn, "ɓe", "village", "stem"),
        "ɓeje": _e(conn, "ɓeje", "chez eux", "stem"),
        "-je": _e(conn, "-je", "PL", "suffix", [Variante("-ge", "allomorphe", environnement="après nasale")]),
        "ka-": _e(conn, "ka-", "NMLZ", "prefix"),
        "=g": _e(conn, "=g", "DEF", "enclitic"),
        "man": _e(conn, "man", "eau"),   # sans type : peut servir de radical
    }
    return projet, ids


# -- analyseur à règles -----------------------------------------------------------------

def test_decomposition_simple_et_allomorphe(lexique_morpho):
    projet, ids = lexique_morpho
    a = Analyseur(projet.conn)
    (d,) = a.analyser("kóje")
    assert d.segmentation == "kó-je" and d.gloses == "arbre-PL"
    (d,) = a.analyser("manGE")      # insensible à la casse, allomorphe « -ge »
    assert d.segmentation == "man-ge" and d.morphemes[1].entree_id == ids["-je"]
    assert d.morphemes[1].environnement == "après nasale"


def test_prefixe_suffixe_clitique(lexique_morpho):
    projet, _ = lexique_morpho
    (d,) = Analyseur(projet.conn).analyser("kakójeg")
    assert d.segmentation == "ka-kó-je-g" and d.gloses == "NMLZ-arbre-PL-DEF"
    assert [m.role for m in d.morphemes] == ["avant", "radical", "apres", "apres"]


def test_ambiguite_les_plus_simples_d_abord(lexique_morpho):
    projet, _ = lexique_morpho
    decs = Analyseur(projet.conn).analyser("ɓeje")
    assert [d.segmentation for d in decs] == ["ɓeje", "ɓe-je"]


def test_rien_n_est_devine(lexique_morpho):
    projet, _ = lexique_morpho
    a = Analyseur(projet.conn)
    assert a.analyser("inconnu") == []
    assert a.analyser("je") == [], "un affixe seul n'est pas un mot"
    assert a.analyser("kója") == [], "aucune forme déclarée ne couvre « a »"


# -- analyses validées et interlinéaire -------------------------------------------------------

def _corpus(projet, paires):
    with projet.conn:
        projet.conn.executemany("INSERT INTO phrases (ref, vernaculaire, analyse, source) VALUES (?,?,?,?)",
                                [(r, v, a, "test") for r, v, a in paires])


def test_statuts_des_mots_dans_une_phrase(lexique_morpho):
    projet, ids = lexique_morpho
    _corpus(projet, [("T001 1", "Kóje ɓeje xyz.", "Des arbres chez eux.")])
    mots = il.phrase(projet, "T001 1", "Kóje ɓeje xyz.")
    assert [(m.texte, m.statut) for m in mots] == [("Kóje", il.PROPOSEE), ("ɓeje", il.AMBIGUE), ("xyz", il.INCONNUE)]
    assert mots[0].analyse.gloses == "arbre-PL" and mots[1].possibilites == 2

    # validation de l'analyse « ɓe-je » : s'applique à toutes les occurrences de la forme
    dec = Analyseur(projet.conn).analyser("ɓeje")[1]
    aid = il.valider(projet.conn, il.Analyse("ɓeje", il.depuis_decomposition("ɓeje", dec).morphemes, "villages"))
    mots = il.phrase(projet, "T001 1", "Kóje ɓeje xyz.")
    assert mots[1].statut == il.VALIDEE and mots[1].analyse.id == aid and mots[1].analyse.glose_mot == "villages"
    assert il.valider(projet.conn, il.Analyse("ɓeje", il.depuis_decomposition("ɓeje", dec).morphemes)) == aid


def test_choix_par_occurrence(lexique_morpho):
    projet, _ = lexique_morpho
    decs = Analyseur(projet.conn).analyser("ɓeje")
    a1 = il.valider(projet.conn, il.depuis_decomposition("ɓeje", decs[0]))
    a2 = il.valider(projet.conn, il.depuis_decomposition("ɓeje", decs[1]))
    texte = "ɓeje ɓeje"
    il.choisir_pour_occurrence(projet.conn, "T001 1", 1, a2)
    mots = il.phrase(projet, "T001 1", texte)
    assert [m.analyse.id for m in mots] == [a1, a2]
    il.choisir_pour_occurrence(projet.conn, "T001 1", 1, None)
    assert [m.analyse.id for m in il.phrase(projet, "T001 1", texte)] == [a1, a1]


def test_validation_refuse_une_segmentation_fausse(lexique_morpho):
    projet, ids = lexique_morpho
    with pytest.raises(ValueError):
        il.valider(projet.conn, il.Analyse("kóje", [il.MorphemeAnalyse("kó", ids["kó"])]))


def test_suppression_entree_garde_l_analyse(lexique_morpho):
    projet, ids = lexique_morpho
    (d,) = Analyseur(projet.conn).analyser("kóje")
    aid = il.valider(projet.conn, il.depuis_decomposition("kóje", d))
    lexique.supprimer_entree(projet.conn, ids["-je"])
    (a,) = il.analyses_de(projet.conn, "kóje")
    assert a.id == aid and a.morphemes[1].entree_id is None and a.segmentation == "kó-je"


# -- inventaire, couverture, concordance ------------------------------------------------------

def test_inventaire_et_couverture(lexique_morpho):
    projet, _ = lexique_morpho
    _corpus(projet, [("T001 1", "kóje man kóje", "a"), ("T001 2", "man xyz", "b")])
    inv = {f.forme: f for f in il.inventaire(projet)}
    assert (inv["kóje"].frequence, inv["kóje"].phrases, inv["kóje"].statut) == (2, 1, il.PROPOSEE)
    assert inv["man"].au_lexique and not inv["kóje"].au_lexique
    assert inv["xyz"].statut == il.INCONNUE
    assert il.couverture(projet) == {il.PROPOSEE: 4, il.INCONNUE: 1}


def test_concordance_modes(lexique_morpho):
    projet, _ = lexique_morpho
    _corpus(projet, [("T001 1", "Le kóje est là, kóje !", "les arbres"), ("T001 2", "makója", "x")])
    assert [(l.gauche, l.cle) for l in il.concordance(projet, "KÓJE")] == [("Le ", "kóje"), ("Le kóje est là, ", "kóje")]
    assert [l.ref for l in il.concordance(projet, "kój", mode="texte")] == ["T001 1", "T001 1", "T001 2"]
    assert len(il.concordance(projet, r"k\w+je", mode="regex")) == 2
    assert il.concordance(projet, "arbres", cote="analyse")[0].traduction == "Le kóje est là, kóje !"
    (d,) = Analyseur(projet.conn).analyser("kóje")
    il.valider(projet.conn, il.depuis_decomposition("kóje", d))
    assert len(il.concordance(projet, "je", mode="morphème")) == 2
    with pytest.raises(ValueError):
        il.concordance(projet, "(", mode="regex")


# -- textes ------------------------------------------------------------------------------------

def test_decoupage_en_phrases():
    assert textes.decouper_en_phrases("Un. Deux ! Trois…\n\nQuatre\ncinq?") == \
        ["Un.", "Deux !", "Trois…", "Quatre", "cinq?"]


def test_import_texte_et_metadonnees(projet):
    meta = textes.Texte(code="", titre="Le lièvre et l'éléphant", genre="conte", locuteur="N. Djimet")
    code = textes.importer_texte(projet, meta, "Phrase un.\nPhrase deux.", "Sentence one.\nSentence two.")
    assert code == "T001"
    assert textes.phrases(projet.conn, code) == [("T001 1", "Phrase un.", "Sentence one."),
                                                 ("T001 2", "Phrase deux.", "Sentence two.")]
    (t,) = textes.lister(projet.conn)
    assert (t.titre, t.genre, t.locuteur, t.phrases) == ("Le lièvre et l'éléphant", "conte", "N. Djimet", 2)
    assert textes.importer_texte(projet, textes.Texte(code=""), "Autre.") == "T002"
    with pytest.raises(ValueError):
        textes.importer_texte(projet, textes.Texte(code=""), "Une. Deux.", "Seulement une.")


def test_tri_des_textes_importes():
    refs = ["T002 1", "T001 10", "T001 2", "T001 1"]
    assert sorted(refs, key=usfm.cle_reference) == ["T001 1", "T001 2", "T001 10", "T002 1"]


def test_liste_des_textes_avec_livres(projet):
    _corpus(projet, [("JOB 1:1", "a", "b"), ("RUT 1:1", "c", "d"), ("MANUEL x", "e", "f")])
    assert [(t.code, t.titre) for t in textes.lister(projet.conn)] == [
        ("RUT", "Ruth"), ("JOB", "Job"), ("MANUEL", "Phrases saisies")]


# -- environnements des allomorphes ---------------------------------------------------------------

from lexicographe.morphologie import Environnement, lire_classes  # noqa: E402


def test_lecture_des_classes():
    assert lire_classes("V = a e i o\nN = m n ŋ ; C = k ŋg") == {
        "V": ["a", "e", "i", "o"], "N": ["m", "n", "ŋ"], "C": ["ŋg", "k"]}


def test_environnements_formels():
    classes = lire_classes("V = a e i o u ; N = m n ŋ ; C = k m n")
    apres_nasale = Environnement("/ N _", classes)
    assert apres_nasale.formel and apres_nasale.accepte("man", "") and not apres_nasale.accepte("ko", "")
    fin_de_mot = Environnement("/ _ #", classes)
    assert fin_de_mot.accepte("ka", "") and not fin_de_mot.accepte("ka", "g")
    debut = Environnement("/ # _ V", classes)
    assert debut.accepte("", "a") and not debut.accepte("k", "a")
    lettre = Environnement("/ o _ ; / [V] _ k", {"V": ["a", "e"]})
    assert lettre.accepte("ko", "") and lettre.accepte("ka", "k") and not lettre.accepte("ka", "")
    for libre in ("après nasale", "/ X _", "/ [inconnue] _", "/ V"):
        e = Environnement(libre, classes)
        assert not e.formel and e.accepte("", ""), "non formel : affiché, jamais appliqué"


def test_analyseur_applique_les_environnements(projet):
    conn = projet.conn
    projet.definir_reglages({"classes_phonologiques": "N = m n ŋ"})
    _e(conn, "kó", "arbre", "stem")
    _e(conn, "man", "eau", "stem")
    _e(conn, "-je", "PL", "suffix", [Variante("-ge", "allomorphe", environnement="/ N _")])
    a = Analyseur(conn)
    assert [d.segmentation for d in a.analyser("mange")] == ["man-ge"]
    assert a.analyser("kóge") == [], "« -ge » n'apparaît qu'après une nasale"
    assert [d.segmentation for d in a.analyser("kóje")] == ["kó-je"]
    # le même allomorphe avec un environnement en texte libre n'est pas filtré
    projet.definir_reglages({"classes_phonologiques": ""})
    assert [d.segmentation for d in Analyseur(conn).analyser("kóge")] == ["kó-ge"], \
        "« / N _ » sans classe N déclarée : non formel, donc pas appliqué"
    e = lexique.lire_entree(conn, lexique.chercher_par_forme(conn, "-je")[0])
    e.variantes[0].environnement = "après nasale"
    lexique.enregistrer_entree(conn, e)
    assert [d.segmentation for d in Analyseur(conn).analyser("kóge")] == ["kó-ge"]


# -- export de l'interlinéaire ------------------------------------------------------------------

def test_export_interlineaire(lexique_morpho, tmp_path):
    projet, _ = lexique_morpho
    code = textes.importer_texte(projet, textes.Texte(code="", titre="Conte"), "Kóje ɓeje xyz.", "Des arbres.")
    with projet.conn:
        projet.conn.execute("UPDATE phrases SET traduction_litterale='arbre-PL', notes='note' WHERE ref=?",
                            (f"{code} 1",))
    chemin = str(tmp_path / "t.txt")
    assert il.exporter_texte_interlineaire(projet, code, "Conte", chemin) == 1
    lignes = open(chemin, encoding="utf-8").read().splitlines()
    i = lignes.index(f"[{code} 1]")
    mots, morphemes, gloses = lignes[i + 1], lignes[i + 2], lignes[i + 3]
    assert mots.startswith("Kóje") and morphemes.startswith("kó-je") and gloses.startswith("arbre-PL")
    assert mots.index("ɓeje") == len("arbre-PL") + 2, "colonne du 1er mot = son élément le plus long + 2"
    assert "Traduction libre : Des arbres." in lignes and "Traduction littérale : arbre-PL" in lignes
    import docx
    chemin = str(tmp_path / "t.docx")
    assert il.exporter_docx_interlineaire(projet, code, "Conte", chemin) == 1
    d = docx.Document(chemin)
    assert d.tables[0].cell(0, 0).text == "Kóje" and d.tables[0].cell(2, 0).text == "arbre-PL"


def test_alignement_avec_tons():
    # « ká » occupe 2 colonnes à l'écran, pas 3 : les marques combinantes ne comptent pas
    assert il._largeur("ká") == 2
