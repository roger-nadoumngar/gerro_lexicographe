import os
import xml.etree.ElementTree as ET

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication  # noqa: E402

from lexicographe import lexique, publication as pub, syllabaire as sy  # noqa: E402
from lexicographe.lexique import Entree, Exemple, Sens  # noqa: E402


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def dico(projet, tmp_path):
    conn = projet.conn
    projet.definir_reglages({"alphabet_vern": "a b k m n ŋ ŋg o"})
    nom = lexique.categorie_par_nom(conn, "nom")
    conn.execute("UPDATE categories SET abreviation='n' WHERE id=?", (nom,))
    conn.commit()
    ids = {}
    for forme, glose, statut in [("kó", "arbre.", "validé"), ("ŋgon", "enfant", "validé"), ("ŋa", "dent", "à vérifier"),
                                 ("ba", "père", "brouillon"), ("kó", "bois", "validé")]:
        ids.setdefault(forme, []).append(lexique.enregistrer_entree(conn, Entree(
            forme=forme, statut=statut, sens=[Sens(glose=glose, categorie_id=nom,
                                                   exemples=[Exemple("kó ɓa", "un arbre", "RUT 1:1")])])))
    audio = tmp_path / "ko.wav"
    audio.write_bytes(b"RIFF....WAVE")
    image = tmp_path / "ko.png"
    image.write_bytes(b"\x89PNG")
    lexique.lier_media(conn, ids["kó"][0], str(audio), "prononciation")
    lexique.lier_media(conn, ids["kó"][0], str(image), "image")
    return projet


def test_contenu(dico):
    c = pub.contenu(dico, pub.Options())
    assert [l for l, _ in c.lettres] == ["K", "Ŋ", "ŊG"], "initiales selon l'alphabet déclaré (ŋg est une lettre)"
    assert c.nb_entrees == 4, "les brouillons ne sont pas publiés par défaut"
    assert [e.vedette for e in c.lettres[0][1]] == ["kó¹", "kó²"]
    assert c.lettres[0][1][0].sens[0].categorie == "n"
    assert ("enfant", ["ŋgon"]) in c.index
    tout = pub.contenu(dico, pub.Options(statuts=("validé", "à vérifier", "brouillon"), exemples=False))
    assert tout.nb_entrees == 5 and not tout.lettres[0][1][0].sens[0].exemples


def test_html_entree_sans_double_point(dico):
    c = pub.contenu(dico, pub.Options())
    h = pub.html_entree(c.lettres[0][1][0])
    assert "<b>kó¹</b>" in h and "<i>n</i>" in h and "arbre." in h and ".." not in h


def test_pdf(app, dico, tmp_path):
    from PySide6.QtPdf import QPdfDocument
    chemin = str(tmp_path / "d.pdf")
    assert pub.exporter_pdf(dico, chemin, pub.Options(titre="Essai", format_page="A5")) == 4
    doc = QPdfDocument()
    doc.load(chemin)
    assert doc.pageCount() >= 3, "page de titre, dictionnaire, index"
    texte = "".join(doc.getAllText(i).text() for i in range(doc.pageCount()))
    for attendu in ("Essai", "ŋgon", "enfant", "Index"):
        assert attendu in texte


def test_docx(dico, tmp_path):
    import docx
    chemin = str(tmp_path / "d.docx")
    pub.exporter_docx(dico, chemin, pub.Options(titre="Essai"))
    d = docx.Document(chemin)
    entrees = [p for p in d.paragraphs if p.style.name == "Entrée"]
    assert entrees[0].runs[0].text == "kó¹" and entrees[0].runs[0].style.name == "Vedette"
    assert d.sections[1]._sectPr.xpath("./w:cols/@w:num") == ["2"]
    assert d.sections[1].header.paragraphs[0]._p.xml.count("STYLEREF") == 2
    assert not any(".." in p.text for p in entrees)


def test_dossier_dictionary_app_builder(dico, tmp_path):
    lift = pub.preparer_dab(dico, str(tmp_path))
    racine = os.path.dirname(lift)
    assert os.path.isfile(os.path.join(racine, "audio", "ko.wav"))
    assert os.path.isfile(os.path.join(racine, "pictures", "ko.png"))
    assert os.path.isfile(os.path.join(racine, "LISEZ-MOI.txt"))
    arbre = ET.parse(lift).getroot()
    assert [m.get("href") for m in arbre.iter("media")] == ["ko.wav"]
    assert [i.get("href") for i in arbre.iter("illustration")] == ["ko.png"]


def test_livret_et_bloom(app, dico, tmp_path):
    from PySide6.QtPdf import QPdfDocument
    with dico.conn:
        dico.conn.executemany("INSERT INTO phrases (ref, vernaculaire, source) VALUES (?, ?, 't')",
                              [("T001 1", "ma ma kó"), ("T001 2", "kó ŋgon")])
    s = sy.Syllabaire(dico)
    lecons = s.lecons([g for g, _ in s.ordre_propose()])
    chemin = str(tmp_path / "livret.pdf")
    assert pub.exporter_livret_syllabaire(dico, lecons, chemin) == len(lecons)
    doc = QPdfDocument()
    doc.load(chemin)
    assert doc.pageCount() == len(lecons) + 1, "page de titre + une page par leçon"
    fichiers = pub.exporter_bloom(dico, lecons, str(tmp_path / "bloom"))
    etapes = open(fichiers[0], encoding="utf-8").read().split()
    assert etapes == [sy.affichage(l.grapheme) for l in lecons]
    assert "ma" in open(fichiers[1], encoding="utf-8").read().split()
